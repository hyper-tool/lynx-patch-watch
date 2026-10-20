"""Excel export: 项目清单 / 项目×公告矩阵 / 问题明细 / 统计.

- git URLs are always masked as ``***`` (internal repo locations are sensitive;
  re-importing ``***`` means “do not modify”);
- cells starting with = + - @ get a leading apostrophe (formula injection);
- statistics are written as numbers, never formulas.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from lynx_patch_watch.engine.orchestrator import project_advisory_status
from lynx_patch_watch.excel.enums import STATUS_ZH

MASKED = "***"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

_STATUS_FILLS = {
    "已修复": PatternFill("solid", fgColor="D6EFE8"),
    "存在漏洞": PatternFill("solid", fgColor="F8D7DA"),
    "冲突待复核": PatternFill("solid", fgColor="FDEBD0"),
    "已声明补丁": PatternFill("solid", fgColor="E8EEF3"),
    "未知": PatternFill("solid", fgColor="EFEFEF"),
    "不适用": PatternFill("solid", fgColor="F4F6F8"),
}

PROJECT_EXPORT_HEADERS = [
    "项目ID", "名称", "类型", "扫描方式", "Git地址", "分支", "本地路径",
    "Magento版本", "版本类型", "已打补丁", "负责人", "环境", "备注", "启用",
    "已修", "未修", "declared", "未知", "补丁维护率", "最长未修天数",
]


def _escape(value: Any) -> Any:
    """Prevent formula injection; pass through non-strings."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _write_header(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col)].width = max(10, len(header) * 2 + 2)
    ws.freeze_panes = "A2"


def build_export(store) -> bytes:
    wb = Workbook()

    projects = store.list_projects()
    advisories = store.list_advisories()

    # ---- 项目清单 ----
    ws = wb.active
    ws.title = "项目清单"
    _write_header(ws, PROJECT_EXPORT_HEADERS)
    for row_idx, project in enumerate(projects, start=2):
        counts = store.count_latest_by_status(project["id"])
        applicable = counts["fixed"] + counts["vulnerable"] + counts["conflict"] + counts["declared"] + counts["unknown"]
        maintenance = round(counts["fixed"] / applicable, 4) if applicable else None
        longest = 0
        for finding in store.list_latest_findings(project["id"]):
            if finding["status"] in ("vulnerable", "conflict", "unknown"):
                from lynx_patch_watch.web.presenters import days_since

                age = days_since(store.first_finding_time(project["id"], finding["item_id"]))
                if age:
                    longest = max(longest, age)
        values = [
            project["id"],
            project.get("name") or "",
            project.get("type") or "",
            project.get("scan_mode") or "",
            MASKED if project.get("git_url") else "",
            project.get("branch") or "",
            project.get("local_path") or "",
            project.get("declared_version") or "",
            project.get("declared_edition") or "",
            ", ".join(project.get("declared_patches") or []),
            ", ".join(project.get("owners") or []),
            project.get("env") or "",
            project.get("notes") or "",
            "是" if project.get("enabled") else "否",
            counts["fixed"],
            counts["vulnerable"] + counts["conflict"] + counts["unknown"],
            counts["declared"],
            counts["unknown"],
            maintenance if maintenance is not None else "—",
            longest,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=_escape(value))

    # ---- 项目×公告矩阵 ----
    ws = wb.create_sheet("项目×公告矩阵")
    headers = ["项目ID", "名称"] + [a["id"] for a in advisories]
    _write_header(ws, headers)
    for row_idx, project in enumerate(projects, start=2):
        statuses = project_advisory_status(store, project["id"])
        ws.cell(row=row_idx, column=1, value=_escape(project["id"]))
        ws.cell(row=row_idx, column=2, value=_escape(project.get("name") or ""))
        for col, adv in enumerate(advisories, start=3):
            status = statuses.get(adv["id"], "unknown")
            label = STATUS_ZH.get(status, status)
            cell = ws.cell(row=row_idx, column=col, value=label)
            fill = _STATUS_FILLS.get(label)
            if fill is not None:
                cell.fill = fill

    # ---- 问题明细 ----
    ws = wb.create_sheet("问题明细")
    _write_header(
        ws,
        [
            "项目ID", "公告", "CVE", "状态", "判定来源", "置信度", "证据摘要",
            "人工状态", "人工备注", "确认是否过期", "首次发现", "修复时间",
        ],
    )
    row_idx = 2
    for project in projects:
        for finding in store.list_latest_findings(project["id"]):
            item = store.get_advisory_item(finding["item_id"])
            review = store.get_manual_review(project["id"], finding["item_id"])
            first = store.first_finding_time(project["id"], finding["item_id"])
            fixed = store.fixed_finding_time(project["id"], finding["item_id"])
            values = [
                project["id"],
                item["advisory_id"] if item else "",
                item.get("cve_id") if item else "",
                STATUS_ZH.get(finding["status"], finding["status"]),
                finding.get("decided_by") or "",
                round(float(finding.get("confidence") or 0), 3),
                (finding.get("evidence_hash") or "")[:16],
                (review or {}).get("status") or "",
                (review or {}).get("note") or "",
                "是" if (review or {}).get("stale") else "否",
                (first or "")[:10],
                (fixed or "")[:10],
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col, value=_escape(value))
            row_idx += 1

    # ---- 统计 ----
    ws = wb.create_sheet("统计")
    _write_header(ws, ["项目ID", "名称", "适用项", "已修", "未修", "declared", "未知", "补丁维护率"])
    for row_idx, project in enumerate(projects, start=2):
        counts = store.count_latest_by_status(project["id"])
        applicable = counts["fixed"] + counts["vulnerable"] + counts["conflict"] + counts["declared"] + counts["unknown"]
        maintenance = round(counts["fixed"] / applicable, 4) if applicable else None
        values = [
            project["id"],
            project.get("name") or "",
            applicable,
            counts["fixed"],
            counts["vulnerable"] + counts["conflict"] + counts["unknown"],
            counts["declared"],
            counts["unknown"],
            maintenance if maintenance is not None else "—",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=_escape(value))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
