"""Excel import template (说明 + 项目 + 人工确认 sheets)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_VERSION = "1.0"

PROJECT_HEADERS = [
    ("项目ID", True),
    ("名称", True),
    ("类型", False),
    ("扫描方式", False),
    ("Git地址", False),
    ("分支", False),
    ("本地路径", False),
    ("Magento版本", False),
    ("版本类型", False),
    ("已打补丁", False),
    ("负责人", False),
    ("环境", False),
    ("备注", False),
    ("启用", False),
]

REVIEW_HEADERS = [
    ("项目ID", True),
    ("公告ID", True),
    ("CVE", False),
    ("状态", True),
    ("备注", False),
    ("确认人", False),
    ("计划日期", False),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_REQUIRED_FONT = Font(color="B91C1C", bold=True)


def build_template() -> bytes:
    wb = Workbook()

    # ---- 说明 ----
    intro = wb.active
    intro.title = "说明"
    intro["A1"] = "lynx-patch-watch 导入模板"
    intro["A1"].font = Font(bold=True, size=14)
    intro["A3"] = f"模板版本: {TEMPLATE_VERSION}"
    intro["A5"] = "工作表说明:"
    intro["A6"] = "  项目: 要扫描/登记的项目清单。按“项目ID”upsert，不会删除表中未出现的项目。"
    intro["A7"] = "  人工确认: 对 (项目, 公告/CVE) 的人工确认结论。CVE 留空表示整条公告。"
    intro["A9"] = "必填列以红色标注；Git 地址中不要写明文凭据（请用 ${ENV_VAR} 占位）。"
    intro["A10"] = "导入规则: 中文枚举值自动映射；是/否/TRUE 均可识别；补丁列表支持全角逗号。"
    intro["A11"] = "导出文件中带 *** 的脱敏值或空 Git 地址，导入时视为“不修改”。"
    intro["A13"] = "本文件为 .xlsx；.xlsm 会被拒绝。"
    intro.column_dimensions["A"].width = 110

    # ---- 项目 ----
    projects = wb.create_sheet("项目")
    for col, (header, required) in enumerate(PROJECT_HEADERS, start=1):
        cell = projects.cell(row=1, column=col, value=header + ("*" if required else ""))
        cell.fill = _HEADER_FILL
        cell.font = _REQUIRED_FONT if required else _HEADER_FONT
        projects.column_dimensions[get_column_letter(col)].width = max(12, len(header) * 2 + 4)
    dv_type = DataValidation(type="list", formula1='"自动识别,magento"', allow_blank=True)
    dv_mode = DataValidation(type="list", formula1='"Git仓库,本地目录,手动登记"', allow_blank=True)
    dv_edition = DataValidation(type="list", formula1='"CE,EE,B2B"', allow_blank=True)
    dv_bool = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    for dv in (dv_type, dv_mode, dv_edition, dv_bool):
        projects.add_data_validation(dv)
    dv_type.add("C2:C500")
    dv_mode.add("D2:D500")
    dv_edition.add("I2:I500")
    dv_bool.add("N2:N500")
    # Format the ID column as text so Excel keeps digit-only IDs verbatim
    # (e.g. "20260101" must not become a number / lose leading zeros).
    for row in range(2, 501):
        projects.cell(row=row, column=1).number_format = "@"
    projects.freeze_panes = "A2"

    # ---- 人工确认 ----
    reviews = wb.create_sheet("人工确认")
    for col, (header, required) in enumerate(REVIEW_HEADERS, start=1):
        cell = reviews.cell(row=1, column=col, value=header + ("*" if required else ""))
        cell.fill = _HEADER_FILL
        cell.font = _REQUIRED_FONT if required else _HEADER_FONT
        reviews.column_dimensions[get_column_letter(col)].width = max(12, len(header) * 2 + 4)
    dv_status = DataValidation(
        type="list",
        formula1='"确认已修,确认存在漏洞,忽略,误报,待定"',
        allow_blank=True,
    )
    reviews.add_data_validation(dv_status)
    dv_status.add("D2:D500")
    # Project ID column as text (see 项目 sheet above).
    for row in range(2, 501):
        reviews.cell(row=row, column=1).number_format = "@"
    reviews.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
